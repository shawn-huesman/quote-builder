from __future__ import annotations

import io
from dataclasses import dataclass, field
from enum import Enum
from typing import Any, Callable, Optional

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from config import PRICING_CONFIG

FONT_NAME = "Arial"
TITLE_FONT = Font(name=FONT_NAME, size=15, bold=True)
HEADER_FONT = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="404040")
BODY_FONT = Font(name=FONT_NAME, size=11)
TOTAL_FONT = Font(name=FONT_NAME, size=11, bold=True)

LOW_MARGIN_FONT = Font(color="9C0006")
LOW_MARGIN_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
LOW_MARGIN_THRESHOLD = 0.20

MED_MARGIN_FONT = Font(color="9C6500")
MED_MARGIN_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
MED_MARGIN_THRESHOLD = 0.30

CURRENCY_FMT = '"$"#,##0.00'
PERCENT_FMT = "0.0%"
MEASUREMENT_FMT = "#,##0.0"

DEFAULT_ROYALTY_RATE = 0.075

PRICE_COLUMN_WIDTH = 16
SUMMARY_LABEL_COLUMN_WIDTH = 22
SUMMARY_SCOPE_COLUMN_WIDTH = 40
TASK_ROOM_COLUMN_WIDTH = 20
TASK_SCOPE_COLUMN_WIDTH = 34

SUMMARY_SECTION_GAP = 2
TASK_TOTAL_ROW_GAP = 2

MEASUREMENT_COLUMNS = (
    ("room_area", "Room Area (sqft)", 16),
    ("room_area_waste", "Room Area (w/Waste) (sqft)", 16),
    ("wall_area", "Wall Area (sqft)", 16),
    ("paintable_sqft", "Paintable Area (sqft)", 18),
    ("perimeter", "Perimeter (ln ft)", 16),
)

MEASUREMENT_MAP = {
    key: (label, width)
    for key, label, width in MEASUREMENT_COLUMNS
}

ALL_MEASUREMENTS = tuple(
    key for key, _, _ in MEASUREMENT_COLUMNS
)

PRICE_OFFSET, MATERIALS_OFFSET, ROYALTY_OFFSET, GROSS_PROFIT_OFFSET, MARGIN_OFFSET = range(5)
PRICE_COLUMN_HEADERS = (
    "HS Price", "Materials", "Royalty", "Gross Profit", "Gross Margin %"
)
NUM_PRICE_COLUMNS = len(PRICE_COLUMN_HEADERS)
SUMMABLE_PRICE_OFFSETS = (
    PRICE_OFFSET, MATERIALS_OFFSET, ROYALTY_OFFSET, GROSS_PROFIT_OFFSET,
)


def fmt_plain(v: Any) -> Any:
    return v if v is not None else ""


def fmt_bool(v: Any) -> str:
    if v is None:
        return ""
    return "Yes" if v else "No"


def fmt_list(v: Any) -> str:
    return ", ".join(v) if v else ""


class FormulaFactory:
    @staticmethod
    def royalty(price_cell: str, rate: float) -> str:
        return f"=PRODUCT({price_cell},{rate})"

    @staticmethod
    def gross_profit(
            price_cell: str, materials_cell: str, royalty_cell: str,
            include_material: bool,
    ) -> str:
        if include_material:
            return f"={price_cell}-{materials_cell}-{royalty_cell}"
        return f"={price_cell}-{royalty_cell}"

    @staticmethod
    def margin(gross_profit_cell: str, price_cell: str) -> str:
        return f"=IFERROR({gross_profit_cell}/{price_cell},0)"

    @staticmethod
    def sumif(criteria_range: str, criteria: str, sum_range: str) -> str:
        return f'=SUMIF({criteria_range},"{criteria}",{sum_range})'

    @staticmethod
    def sum(value_range: str) -> str:
        return f"=SUM({value_range})"


@dataclass(frozen=True)
class PriceColumnLayout:
    start_col: int

    @property
    def price(self) -> int:
        return self.start_col + PRICE_OFFSET

    @property
    def materials(self) -> int:
        return self.start_col + MATERIALS_OFFSET

    @property
    def royalty(self) -> int:
        return self.start_col + ROYALTY_OFFSET

    @property
    def gross_profit(self) -> int:
        return self.start_col + GROSS_PROFIT_OFFSET

    @property
    def margin(self) -> int:
        return self.start_col + MARGIN_OFFSET

    @staticmethod
    def cell_ref(col: int, row: int) -> str:
        return f"{get_column_letter(col)}{row}"


class TaskType(str, Enum):
    INTERIOR_PAINTING = "Interior_Painting"
    EXTERIOR_PAINTING = "Exterior_Painting"
    CABINET_PAINTING = "Cabinet_Painting"
    CARPET_INSTALLATION = "Carpet_Installation"
    FLOORING_INSTALLATION = "Flooring_Installation"
    CARPET_CLEANING = "Carpet_Cleaning"
    CARPET_REMOVAL = "Carpet_Removal"
    HOME_CLEAR_OUTS = "Home_Clear_Outs"
    LIGHT_HANDYMAN_WORK = "Light_Handyman_Work"
    MOVE_OUT_CLEANING = "Move_Out_Cleaning"
    LANDSCAPE_CLEAN_UP = "Landscape_Clean_Up"
    PRESSURE_WASHING = "Pressure_Washing"
    WINDOW_WASHING = "Window_Washing"


@dataclass(frozen=True)
class ColumnSpec:
    key: str
    header: str
    formatter: Callable[[Any], Any] = fmt_plain
    width: int = 14


@dataclass(frozen=True)
class TaskTypeSpec:
    display_name: str
    detail_columns: tuple[ColumnSpec, ...] = ()
    include_material_in_profit: bool = True


TASK_TYPE_REGISTRY: dict[TaskType, TaskTypeSpec] = {
    TaskType.INTERIOR_PAINTING: TaskTypeSpec(
        "Interior Painting",
        (
            ColumnSpec("sub_tasks", "Sub-Tasks", fmt_list, 30),
            ColumnSpec("num_doors", "Doors"),
            ColumnSpec("has_closet", "Closet?", fmt_bool),
            ColumnSpec("patches_small", "Small Patches"),
            ColumnSpec("patches_medium", "Medium Patches"),
            ColumnSpec("patches_large", "Large Patches"),
        ),
    ),
    TaskType.EXTERIOR_PAINTING: TaskTypeSpec(
        "Exterior Painting",
        (
            ColumnSpec("num_sides", "Sides"),
            ColumnSpec("stories", "Stories"),
            ColumnSpec("substrate", "Substrate"),
            ColumnSpec("deck_size", "Deck Size"),
        ),
    ),
    TaskType.CABINET_PAINTING: TaskTypeSpec(
        "Cabinet Painting",
        (
            ColumnSpec("num_doors", "Doors"),
            ColumnSpec("single_vanities", "Single Vanities"),
            ColumnSpec("double_vanities", "Double Vanities"),
        ),
    ),
    TaskType.CARPET_INSTALLATION: TaskTypeSpec(
        "Carpet Installation",
        (
            ColumnSpec("carpet_tier", "Tier"),
            ColumnSpec("stair_steps", "Stair Steps"),
            ColumnSpec("kilz_required", "Kilz?", fmt_bool),
            ColumnSpec("demo_required", "Demo?", fmt_bool),
            ColumnSpec("ceramic_demo_required", "Ceramic Demo?", fmt_bool),
            ColumnSpec("furniture_move", "Furniture Move?", fmt_bool),
        ),
    ),
    TaskType.FLOORING_INSTALLATION: TaskTypeSpec(
        "Flooring Installation",
        (
            ColumnSpec("material", "Material"),
            ColumnSpec("transitions", "Transitions"),
            ColumnSpec("kilz_required", "Kilz?", fmt_bool),
            ColumnSpec("demo_required", "Demo?", fmt_bool),
            ColumnSpec("ceramic_demo_required", "Ceramic Demo?", fmt_bool),
            ColumnSpec("furniture_move", "Furniture Move?", fmt_bool),
            ColumnSpec("appliances_to_move", "Appliances to Move"),
            ColumnSpec("toilets_to_reset", "Toilets to Reset"),
        ),
    ),
    TaskType.CARPET_CLEANING: TaskTypeSpec(
        "Carpet Cleaning",
        (
            ColumnSpec("room_count", "Rooms"),
            ColumnSpec("pet_treatment_required", "Pet Treatment?", fmt_bool),
        ),
    ),
    TaskType.CARPET_REMOVAL: TaskTypeSpec(
        "Carpet Removal",
        (ColumnSpec("requested", "Requested?", fmt_bool),),
    ),
    TaskType.HOME_CLEAR_OUTS: TaskTypeSpec(
        "Home Clear Outs",
        (
            ColumnSpec("box_truck_loads", "Box Truck Loads"),
            ColumnSpec("dumpster_20yd_count", "20yd Dumpsters"),
            ColumnSpec("dumpster_30yd_count", "30yd Dumpsters"),
            ColumnSpec("dumpster_40yd_count", "40yd Dumpsters"),
            ColumnSpec("tire_count", "Tires"),
            ColumnSpec("mattress_count", "Mattresses"),
            ColumnSpec("refrigerator_count", "Refrigerators"),
        ),
    ),
    TaskType.LIGHT_HANDYMAN_WORK: TaskTypeSpec(
        "Light Handyman Work",
        (
            ColumnSpec("task_count", "Tasks"),
            ColumnSpec("cabinet_door_pull_amount", "Cabinet Pulls"),
            ColumnSpec("light_fixture_amount", "Light Fixtures"),
            ColumnSpec("ceiling_fan_amount", "Ceiling Fans"),
            ColumnSpec("mirror_amount", "Mirrors"),
        ),
    ),
    TaskType.MOVE_OUT_CLEANING: TaskTypeSpec(
        "Move Out Cleaning",
        (ColumnSpec("status", "Condition"),),
    ),
    TaskType.LANDSCAPE_CLEAN_UP: TaskTypeSpec(
        "Landscape Clean-Up",
        (
            ColumnSpec("total_man_days", "Man Days"),
            ColumnSpec("num_beds_mulched", "Beds Mulched"),
            ColumnSpec("num_bags_mulch", "Mulch Bags"),
        ),
    ),
    TaskType.PRESSURE_WASHING: TaskTypeSpec(
        "Pressure Washing",
        (
            ColumnSpec("driveway_requested", "Driveway?", fmt_bool),
            ColumnSpec("porch_requested", "Porch?", fmt_bool),
            ColumnSpec("sidewalk_requested", "Sidewalk?", fmt_bool),
            ColumnSpec("deck_requested", "Deck?", fmt_bool),
            ColumnSpec("exterior_siding_requested", "Siding?", fmt_bool),
        ),
    ),
    TaskType.WINDOW_WASHING: TaskTypeSpec(
        "Window Washing",
        (ColumnSpec("pane_count", "Panes"),),
    ),
}


def _spec_for(task_type: str, sample_details: dict) -> TaskTypeSpec:
    if task_type in TASK_TYPE_REGISTRY:
        return TASK_TYPE_REGISTRY[task_type]

    fallback_columns = tuple(
        ColumnSpec(key, key.replace("_", " ").title())
        for key in sample_details
        if key != "service_description"
    )
    return TaskTypeSpec(task_type.replace("_", " "), fallback_columns)


def _measurement_headers() -> list[str]:
    return [label for _, label, _ in MEASUREMENT_COLUMNS]


@dataclass
class PricedTaskRecord:
    room_name: str
    task_type: str
    details: dict = field(default_factory=dict)
    price: float = 0.0
    payout: float = 0.0
    material_cost: float = 0.0
    scope: Optional[str] = None
    room_area: Optional[float] = None
    room_area_waste: Optional[float] = None
    wall_area: Optional[float] = None
    paintable_sqft: Optional[float] = None
    perimeter: Optional[float] = None

    @property
    def materials(self) -> float:
        return self.payout + self.material_cost

    def resolved_scope(self) -> str:
        if self.scope:
            return self.scope
        return self.details.get("service_description", "")


@dataclass
class SummaryTableResult:
    table_start: int
    first_data_row: int
    last_data_row: int
    price_start_col: int
    num_headers: int


@dataclass
class SubtotalsContext:
    ordered_types: list
    first_data_row: int
    last_data_row: int
    price_start_col: int


@dataclass
class TaskTableResult:
    header_row: int
    first_data_row: int
    last_data_row: int
    price_col: int
    num_headers: int


class _BaseSheetWriter:

    def __init__(self, royalty_rate: float):
        self.royalty_rate = royalty_rate

    def _write_header_row(self, ws: Worksheet, row: int, headers: list[str]) -> None:
        for i, h in enumerate(headers, start=1):
            c = ws.cell(row=row, column=i, value=h)
            c.font = HEADER_FONT
            c.fill = HEADER_FILL

    def _autofit_and_filter(
        self, ws: Worksheet, header_row: int, last_data_row: int, num_cols: int
    ) -> None:
        if last_data_row < header_row:
            return
        last_col_letter = get_column_letter(num_cols)
        ws.auto_filter.ref = f"A{header_row}:{last_col_letter}{last_data_row}"


    def _apply_margin_formatting(
            self, ws: Worksheet, first_row: int, last_row: int, col: int
    ) -> None:
        if last_row < first_row:
            return
        letter = get_column_letter(col)
        cell_range = f"{letter}{first_row}:{letter}{last_row}"

        low_rule = CellIsRule(
            operator="lessThan",
            formula=[str(LOW_MARGIN_THRESHOLD)],
            font=LOW_MARGIN_FONT,
            fill=LOW_MARGIN_FILL,
        )

        med_rule = CellIsRule(
            operator="between",
            formula=[str(LOW_MARGIN_THRESHOLD), str(MED_MARGIN_THRESHOLD)],
            font=MED_MARGIN_FONT,
            fill=MED_MARGIN_FILL,
        )

        ws.conditional_formatting.add(cell_range, low_rule)
        ws.conditional_formatting.add(cell_range, med_rule)

    def _write_measurement_columns(
        self, ws: Worksheet, row: int, start_col: int, rec: PricedTaskRecord
    ) -> int:
        col = start_col
        for attr, _, _ in MEASUREMENT_COLUMNS:
            value = getattr(rec, attr, None)
            c = ws.cell(row=row, column=col, value=value if value is not None else "")
            if value is not None:
                c.number_format = MEASUREMENT_FMT
            c.font = BODY_FONT
            col += 1
        return col

    def _write_price_columns(
        self, ws: Worksheet, row: int, start_col: int,
        rec: PricedTaskRecord, spec: TaskTypeSpec,
    ) -> None:
        layout = PriceColumnLayout(start_col)
        price_ref = layout.cell_ref(layout.price, row)
        materials_ref = layout.cell_ref(layout.materials, row)
        royalty_ref = layout.cell_ref(layout.royalty, row)
        gross_profit_ref = layout.cell_ref(layout.gross_profit, row)

        p = ws.cell(row=row, column=layout.price, value=rec.price)
        p.number_format = CURRENCY_FMT
        p.font = BODY_FONT

        mat = ws.cell(row=row, column=layout.materials, value=rec.materials)
        mat.number_format = CURRENCY_FMT
        mat.font = BODY_FONT

        roy = ws.cell(
            row=row, column=layout.royalty,
            value=FormulaFactory.royalty(price_ref, self.royalty_rate),
        )
        roy.number_format = CURRENCY_FMT
        roy.font = BODY_FONT

        gp = ws.cell(
            row=row, column=layout.gross_profit,
            value=FormulaFactory.gross_profit(
                price_ref, materials_ref, royalty_ref,
                spec.include_material_in_profit,
            ),
        )
        gp.number_format = CURRENCY_FMT
        gp.font = BODY_FONT

        margin = ws.cell(
            row=row, column=layout.margin,
            value=FormulaFactory.margin(gross_profit_ref, price_ref),
        )
        margin.number_format = PERCENT_FMT
        margin.font = BODY_FONT


class SummarySheetWriter(_BaseSheetWriter):

    def write(
        self,
        wb: Workbook,
        records: list[PricedTaskRecord],
        ordered_types: list,
        job_meta: Optional[dict],
    ) -> None:
        ws = wb.create_sheet("Summary")

        r = self._write_title_and_meta(ws, job_meta)

        table = self._write_line_item_table(ws, r, records)
        self._autofit_and_filter(ws, table.table_start, table.last_data_row, table.num_headers)
        margin_col = table.price_start_col + MARGIN_OFFSET
        self._apply_margin_formatting(ws, table.first_data_row, table.last_data_row, col=margin_col)

        r = table.last_data_row + 1 + SUMMARY_SECTION_GAP

        ctx = SubtotalsContext(
            ordered_types=ordered_types,
            first_data_row=table.first_data_row,
            last_data_row=table.last_data_row,
            price_start_col=table.price_start_col,
        )
        r, subtotal_first_row, subtotal_last_row = self._write_subtotals(ws, r, ctx)

        self._write_grand_total(ws, r, subtotal_first_row, subtotal_last_row)

        self._set_column_widths(ws)
        ws.freeze_panes = f"A{table.first_data_row}"

    def _write_title_and_meta(self, ws: Worksheet, job_meta: Optional[dict]) -> int:
        r = 1
        ws.cell(row=r, column=1, value="Quote Summary").font = TITLE_FONT
        r += 2

        if job_meta:
            for label, value in job_meta.items():
                ws.cell(row=r, column=1, value=f"{label}:").font = TOTAL_FONT
                ws.cell(row=r, column=2, value=value).font = BODY_FONT
                r += 1
            r += 1

        return r

    def _write_line_item_table(
        self, ws: Worksheet, r: int, records: list[PricedTaskRecord]
    ) -> SummaryTableResult:
        table_start = r
        price_start_col = 3 + len(MEASUREMENT_COLUMNS) + 1
        headers = (
            ["Room", "Task Type"] + _measurement_headers() + ["Scope"] + list(PRICE_COLUMN_HEADERS)
        )
        self._write_header_row(ws, r, headers)
        r += 1
        first_data_row = r

        for rec in records:
            spec = _spec_for(rec.task_type, rec.details)
            ws.cell(row=r, column=1, value=rec.room_name).font = BODY_FONT
            ws.cell(row=r, column=2, value=spec.display_name).font = BODY_FONT
            scope_col = self._write_measurement_columns(ws, r, 3, rec)
            ws.cell(row=r, column=scope_col, value=rec.resolved_scope()).font = BODY_FONT
            self._write_price_columns(ws, r, scope_col + 1, rec, spec)
            r += 1

        last_data_row = r - 1
        return SummaryTableResult(table_start, first_data_row, last_data_row, price_start_col, len(headers))

    def _write_subtotals(
        self, ws: Worksheet, r: int, ctx: SubtotalsContext,
    ) -> tuple[int, int, int]:
        ws.cell(row=r, column=1, value="Totals by Task Type").font = TOTAL_FONT
        r += 1
        self._write_header_row(
            ws, r, ["Task Type"] + [PRICE_COLUMN_HEADERS[i] for i in SUMMABLE_PRICE_OFFSETS],
        )
        r += 1
        subtotal_first_row = r
        type_col_range = f"B{ctx.first_data_row}:B{ctx.last_data_row}"
        source_cols = [
            get_column_letter(ctx.price_start_col + offset) for offset in SUMMABLE_PRICE_OFFSETS
        ]

        for task_type in ctx.ordered_types:
            spec = TASK_TYPE_REGISTRY.get(task_type) or _spec_for(task_type, {})
            ws.cell(row=r, column=1, value=spec.display_name).font = BODY_FONT
            for i, source_letter in enumerate(source_cols):
                dest_col = 2 + i
                sum_range = f"{source_letter}{ctx.first_data_row}:{source_letter}{ctx.last_data_row}"
                formula = FormulaFactory.sumif(type_col_range, spec.display_name, sum_range)
                c = ws.cell(row=r, column=dest_col, value=formula)
                c.number_format = CURRENCY_FMT
                c.font = BODY_FONT
            r += 1

        subtotal_last_row = r - 1
        r += 1
        return r, subtotal_first_row, subtotal_last_row

    def _write_grand_total(
        self, ws: Worksheet, r: int, subtotal_first_row: int, subtotal_last_row: int
    ) -> None:
        ws.cell(row=r, column=1, value="Grand Total").font = TOTAL_FONT
        for col in range(2, 2 + len(SUMMABLE_PRICE_OFFSETS)):
            letter = get_column_letter(col)
            value_range = f"{letter}{subtotal_first_row}:{letter}{subtotal_last_row}"
            c = ws.cell(row=r, column=col, value=FormulaFactory.sum(value_range))
            c.number_format = CURRENCY_FMT
            c.font = TOTAL_FONT

    def _set_column_widths(self, ws: Worksheet) -> None:
        ws.column_dimensions["A"].width = SUMMARY_LABEL_COLUMN_WIDTH
        ws.column_dimensions["B"].width = SUMMARY_LABEL_COLUMN_WIDTH
        col = 3
        for _, _, width in MEASUREMENT_COLUMNS:
            ws.column_dimensions[get_column_letter(col)].width = width
            col += 1
        ws.column_dimensions[get_column_letter(col)].width = SUMMARY_SCOPE_COLUMN_WIDTH
        for offset in range(1, NUM_PRICE_COLUMNS + 1):
            ws.column_dimensions[get_column_letter(col + offset)].width = PRICE_COLUMN_WIDTH


class TaskSheetWriter(_BaseSheetWriter):

    def write(self, wb: Workbook, task_type, records: list[PricedTaskRecord]) -> None:
        spec = _spec_for(task_type, records[0].details if records else {})
        ws = wb.create_sheet(spec.display_name[:31])

        r = 1
        ws.cell(row=r, column=1, value=spec.display_name).font = TITLE_FONT
        r += 2

        table = self._write_table(ws, r, spec, records)

        if table.last_data_row >= table.first_data_row:
            self._write_total_row(
                ws, table.last_data_row + TASK_TOTAL_ROW_GAP,
                table.first_data_row, table.last_data_row, table.price_col,
            )

        self._apply_margin_formatting(ws, table.first_data_row, table.last_data_row, col=table.price_col + MARGIN_OFFSET)
        self._autofit_and_filter(ws, table.header_row, table.last_data_row, table.num_headers)

        self._set_column_widths(ws, spec)
        ws.freeze_panes = f"A{table.first_data_row}"

    def _write_table(
        self, ws: Worksheet, r: int, spec: TaskTypeSpec, records: list[PricedTaskRecord]
    ) -> TaskTableResult:
        headers = (
            ["Room"] + _measurement_headers() + [c.header for c in spec.detail_columns]
            + ["Scope"] + list(PRICE_COLUMN_HEADERS)
        )
        self._write_header_row(ws, r, headers)
        header_row = r
        r += 1
        first_data_row = r

        price_col = 2 + len(MEASUREMENT_COLUMNS) + len(spec.detail_columns) + 1

        for rec in records:
            ws.cell(row=r, column=1, value=rec.room_name).font = BODY_FONT
            col = self._write_measurement_columns(ws, r, 2, rec)
            for detail_col in spec.detail_columns:
                value = detail_col.formatter(rec.details.get(detail_col.key))
                ws.cell(row=r, column=col, value=value).font = BODY_FONT
                col += 1
            ws.cell(row=r, column=col, value=rec.resolved_scope()).font = BODY_FONT
            col += 1
            self._write_price_columns(ws, r, col, rec, spec)
            r += 1

        last_data_row = r - 1
        return TaskTableResult(header_row, first_data_row, last_data_row, price_col, len(headers))

    def _write_total_row(
        self, ws: Worksheet, r: int, first_data_row: int, last_data_row: int, price_col: int
    ) -> None:
        ws.cell(row=r, column=1, value="Total").font = TOTAL_FONT
        for offset in SUMMABLE_PRICE_OFFSETS:
            col = price_col + offset
            letter = get_column_letter(col)
            value_range = f"{letter}{first_data_row}:{letter}{last_data_row}"
            c = ws.cell(row=r, column=col, value=FormulaFactory.sum(value_range))
            c.number_format = CURRENCY_FMT
            c.font = TOTAL_FONT

        margin_col = price_col + MARGIN_OFFSET
        price_ref = PriceColumnLayout.cell_ref(price_col, r)
        gp_ref = PriceColumnLayout.cell_ref(price_col + GROSS_PROFIT_OFFSET, r)
        c = ws.cell(row=r, column=margin_col, value=FormulaFactory.margin(gp_ref, price_ref))
        c.number_format = PERCENT_FMT
        c.font = TOTAL_FONT

    def _set_column_widths(self, ws: Worksheet, spec: TaskTypeSpec) -> None:
        ws.column_dimensions["A"].width = TASK_ROOM_COLUMN_WIDTH
        col = 2
        for _, _, width in MEASUREMENT_COLUMNS:
            ws.column_dimensions[get_column_letter(col)].width = width
            col += 1
        for detail_col in spec.detail_columns:
            ws.column_dimensions[get_column_letter(col)].width = detail_col.width
            col += 1
        ws.column_dimensions[get_column_letter(col)].width = TASK_SCOPE_COLUMN_WIDTH
        for offset in range(1, NUM_PRICE_COLUMNS + 1):
            ws.column_dimensions[get_column_letter(col + offset)].width = PRICE_COLUMN_WIDTH


class QuoteExcelBuilder:
    def __init__(self, royalty_rate: float = DEFAULT_ROYALTY_RATE):
        self.royalty_rate = royalty_rate
        self._summary_writer = SummarySheetWriter(royalty_rate)
        self._task_writer = TaskSheetWriter(royalty_rate)

    def build(
        self,
        records: list[PricedTaskRecord],
        job_meta: Optional[dict] = None,
    ) -> Workbook:
        wb = Workbook()
        wb.remove(wb.active)

        by_type: dict[str, list[PricedTaskRecord]] = {}
        for rec in records:
            by_type.setdefault(rec.task_type, []).append(rec)

        ordered_types = [t for t in TASK_TYPE_REGISTRY if t in by_type]
        ordered_types += [t for t in by_type if t not in TASK_TYPE_REGISTRY]

        self._summary_writer.write(wb, records, ordered_types, job_meta)
        for task_type in ordered_types:
            self._task_writer.write(wb, task_type, by_type[task_type])

        return wb

    def save(
        self,
        path: str,
        records: list[PricedTaskRecord],
        job_meta: Optional[dict] = None,
    ) -> None:
        self.build(records, job_meta).save(path)

    def build_from_job(
        self,
        job,
        only_priced: bool = True,
    ) -> Workbook:
        records = price_records_from_job(
            job,
            only_priced=only_priced,
        )
        return self.build(records, job_meta_from_job(job))

    def build_to_buffer(
            self,
            job,
            only_priced: bool = True,
    ) -> io.BytesIO:
        wb = self.build_from_job(job, only_priced=only_priced)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    def save_from_job(
        self,
        path: str,
        job,
        only_priced: bool = True,
    ) -> None:
        self.build_from_job(
            job, only_priced
        ).save(path)


def get_attr_if_type(container, attr_name: str, task_type: str, allowed_types: set) -> Any:
    return getattr(container, attr_name, None) if task_type in allowed_types else None


def price_records_from_job(
    job,
    engine=None,
    only_priced: bool = True,
) -> list[PricedTaskRecord]:

    def make_record(task, room_name: str, container) -> PricedTaskRecord:
        if engine is not None:
            task.price = engine.price_task(task.task_type, task.details, container)

        waste_multipler = 1
        if task.task_type == "Flooring_Installation":
            waste_multipler = PRICING_CONFIG["Flooring_Installation"].get("waste_factor_multiplier")
        elif task.task_type == "Carpet_Installation":
            waste_multipler = PRICING_CONFIG["Carpet_Installation"].get("waste_factor_multiplier")

        room_area_waste = None
        room_area = get_attr_if_type(container, "ceiling_area", task.task_type, {"Flooring_Installation",
                                                                                    "Carpet_Installation"})
        if room_area:
            room_area_waste = room_area * waste_multipler

        wall_area = get_attr_if_type(container, "wall_area", task.task_type, {"Interior_Painting"})
        paintable_sqft = get_attr_if_type(container, "paintable_sqft", task.task_type, {"Interior_Painting"})
        perimeter = get_attr_if_type(container, "perimeter", task.task_type, {"Interior_Painting"})

        return PricedTaskRecord(
            room_name=room_name,
            task_type=task.task_type,
            details=task.details,
            price=task.price,
            payout=task.payout_price,
            material_cost=task.material_price,
            scope=task.service_description or None,
            room_area=room_area,
            room_area_waste=room_area_waste,
            wall_area=wall_area,
            paintable_sqft=paintable_sqft,
            perimeter=perimeter,
        )

    records: list[PricedTaskRecord] = []
    for space in job.spaces:
        for room in space.rooms:
            for task in room.tasks:
                record = make_record(task, room.label, room)
                if only_priced and record.price <= 0:
                    continue
                records.append(record)
        for task in space.tasks:
            record = make_record(task, space.name, space)
            if only_priced and record.price <= 0:
                continue
            records.append(record)

    return records


def job_meta_from_job(job) -> dict:
    meta = {
        "Address": job.address
    }
    if getattr(job, "owner_email", None):
        meta["Owner Email"] = job.owner_email
    if getattr(job, "project_notes", None):
        meta["Project Notes"] = job.project_notes
    return meta