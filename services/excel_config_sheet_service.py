import io
import re
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from concurrent.futures import ThreadPoolExecutor

import config

try:
    from google.oauth2 import service_account
    from googleapiclient.discovery import build
    from googleapiclient.http import MediaIoBaseDownload, MediaIoBaseUpload

    HAS_GOOGLE_DRIVE = True
except ImportError:
    HAS_GOOGLE_DRIVE = False


class MarketConfigManager:
    HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    SECTION_FILL = PatternFill(start_color="374151", end_color="374151", fill_type="solid")
    UNLOCKED_FILL = PatternFill(start_color="F0F9FF", end_color="F0F9FF", fill_type="solid")
    LOCKED_FILL = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

    WHITE_HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    SECTION_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    KEY_FONT = Font(name="Calibri", size=11, color="111827")
    EDITABLE_FONT = Font(name="Calibri", size=11, bold=True, color="0369A1")

    THIN_BORDER = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )

    @staticmethod
    def flatten_dict(d, parent_key=''):
        items = []
        for k, v in d.items():
            new_key = f"{parent_key}.{k}" if parent_key else str(k)
            if isinstance(v, dict):
                items.extend(MarketConfigManager.flatten_dict(v, new_key).items())
            else:
                items.append((new_key, v))
        return dict(items)

    @staticmethod
    def deep_merge(default_dict, excel_dict):
        merged = default_dict.copy()
        for k, v in excel_dict.items():
            if k in merged:
                if isinstance(v, dict) and isinstance(merged[k], dict):
                    merged[k] = MarketConfigManager.deep_merge(merged[k], v)
                else:
                    merged[k] = v
            else:
                merged[k] = v
        return merged

    @staticmethod
    def unflatten_dict(flat_dict):
        result = {}
        for key, value in flat_dict.items():
            parts = str(key).split('.')
            d = result
            for part in parts[:-1]:
                if part.isdigit():
                    part = int(part)
                if part not in d:
                    d[part] = {}
                d = d[part]

            last_part = parts[-1]
            if last_part.isdigit():
                last_part = int(last_part)

            d[last_part] = float(value)
        return result

    @classmethod
    def get_default_configs(cls):
        return {
            "Pricing Config": getattr(config, "PRICING_CONFIG", {}),
            "Material Pricing Config": getattr(config, "MATERIAL_PRICING_CONFIG", {}),
            "Payout Pricing Config": getattr(config, "PAYOUT_PRICING_CONFIG", {})
        }

    @classmethod
    def create_excel_workbook(cls, configs_dict=None):
        if configs_dict is None:
            configs_dict = cls.get_default_configs()

        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        for sheet_name, config_data in configs_dict.items():
            ws = wb.create_sheet(title=sheet_name)
            ws.views.sheetView[0].showGridLines = True
            ws.freeze_panes = 'A2'

            headers = ["Key Path / Parameter ID", "Value (Editable)", "Category / Section"]
            for col_num, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_num, value=header)
                cell.font = cls.WHITE_HEADER_FONT
                cell.fill = cls.HEADER_FILL
                cell.alignment = Alignment(horizontal="center" if col_num == 2 else "left", vertical="center")
            ws.row_dimensions[1].height = 28

            row_idx = 2
            for key, val in config_data.items():
                if isinstance(val, dict):
                    ws.cell(row=row_idx, column=1, value=f"[{key}]").font = cls.SECTION_FONT
                    for col_idx in range(1, 4):
                        ws.cell(row=row_idx, column=col_idx).fill = cls.SECTION_FILL
                    ws.row_dimensions[row_idx].height = 22
                    row_idx += 1

                    flat_sub = cls.flatten_dict(val)
                    for sub_key, sub_val in flat_sub.items():
                        full_path = f"{key}.{sub_key}"

                        c1 = ws.cell(row=row_idx, column=1, value=full_path)
                        c1.font, c1.fill, c1.border = cls.KEY_FONT, cls.LOCKED_FILL, cls.THIN_BORDER

                        c2 = ws.cell(row=row_idx, column=2, value=float(sub_val))
                        c2.font, c2.fill, c2.border = cls.EDITABLE_FONT, cls.UNLOCKED_FILL, cls.THIN_BORDER
                        c2.number_format = '#,##0.00'
                        c2.alignment = Alignment(horizontal="right")
                        c2.protection = openpyxl.styles.Protection(locked=False)

                        c3 = ws.cell(row=row_idx, column=3, value=key)
                        c3.font, c3.fill, c3.border = cls.KEY_FONT, cls.LOCKED_FILL, cls.THIN_BORDER

                        ws.row_dimensions[row_idx].height = 20
                        row_idx += 1
                else:
                    c1 = ws.cell(row=row_idx, column=1, value=key)
                    c1.font, c1.fill, c1.border = cls.KEY_FONT, cls.LOCKED_FILL, cls.THIN_BORDER

                    c2 = ws.cell(row=row_idx, column=2, value=float(val))
                    c2.font, c2.fill, c2.border = cls.EDITABLE_FONT, cls.UNLOCKED_FILL, cls.THIN_BORDER
                    c2.number_format = '#,##0.00'
                    c2.alignment = Alignment(horizontal="right")
                    c2.protection = openpyxl.styles.Protection(locked=False)

                    c3 = ws.cell(row=row_idx, column=3, value="General Global")
                    c3.font, c3.fill, c3.border = cls.KEY_FONT, cls.LOCKED_FILL, cls.THIN_BORDER

                    ws.row_dimensions[row_idx].height = 20
                    row_idx += 1

            dv = DataValidation(
                type="decimal", operator="greaterThanOrEqual", formula1=-999999999, allow_blank=False,
                errorTitle="Invalid Entry", error="Entry must be a number (float/decimal). Text/blanks are not allowed."
            )
            ws.add_data_validation(dv)
            dv.add(f"B2:B{row_idx}")

            ws.protection.sheet = True
            ws.protection.enable()

            ws.column_dimensions['A'].width = 45
            ws.column_dimensions['B'].width = 20
            ws.column_dimensions['C'].width = 25

        return wb

    @classmethod
    def parse_workbook(cls, wb):
        parsed_configs = {}
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            flat_data = {}
            for row in ws.iter_rows(min_row=2, values_only=True):
                key_path, val = row[0], row[1]
                if not key_path or str(key_path).startswith("[") or val is None:
                    continue
                try:
                    flat_data[str(key_path)] = float(val)
                except (ValueError, TypeError):
                    continue
            parsed_configs[sheet_name] = cls.unflatten_dict(flat_data)

        defaults = cls.get_default_configs()
        return cls.deep_merge(defaults, parsed_configs)

    @classmethod
    def save_to_file(cls, configs_dict=None, filepath="market_pricing_config.xlsx"):
        if configs_dict is None:
            configs_dict = cls.get_default_configs()
        wb = cls.create_excel_workbook(configs_dict)
        wb.save(filepath)

    @classmethod
    def load_from_file(cls, filepath="market_pricing_config.xlsx"):
        wb = openpyxl.load_workbook(filepath, data_only=True)
        return cls.parse_workbook(wb)

    @classmethod
    def list_workspace_files_in_folder(cls, excel_key, delegate_user, folder_id: str) -> dict:
        scopes = [
            'https://www.googleapis.com/auth/drive',
            'https://www.googleapis.com/auth/spreadsheets'
        ]
        credentials = service_account.Credentials.from_service_account_info(excel_key, scopes=scopes)
        if delegate_user:
            credentials = credentials.with_subject(delegate_user)

        drive_service = build('drive', 'v3', credentials=credentials)

        query = (
            f"'{folder_id}' in parents and "
            f"mimeType = 'application/vnd.google-apps.spreadsheet' and "
            f"trashed = false"
        )

        print(drive_service, credentials)

        results = drive_service.files().list(
            q=query,
            supportsAllDrives=True,
            includeItemsFromAllDrives=True,
            fields="files(id, name)"
        ).execute()

        files = results.get("files", [])
        workspace_map = {}

        for file in files:
            name = file.get("name", "")
            match = re.search(r"\[(.*?)\]$", name)
            if match:
                ws_id = match.group(1)
                workspace_map[ws_id] = {
                    "file_id": file.get("id"),
                    "name": name
                }
        return workspace_map

    @classmethod
    def fetch_and_parse_from_drive(cls, excel_key, delegate_user, file_id):
        scopes = [
            'https://www.googleapis.com/auth/drive',
            'https://www.googleapis.com/auth/spreadsheets'
        ]
        credentials = service_account.Credentials.from_service_account_info(excel_key, scopes=scopes)
        if delegate_user:
            credentials = credentials.with_subject(delegate_user)

        drive_service = build('drive', 'v3', credentials=credentials)

        request = drive_service.files().export_media(
            fileId=file_id,
            mimeType='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        file_stream = io.BytesIO()
        downloader = MediaIoBaseDownload(file_stream, request)

        done = False
        while not done:
            _, done = downloader.next_chunk()

        file_stream.seek(0)
        wb = openpyxl.load_workbook(file_stream, data_only=True)
        return cls.parse_workbook(wb)

    @classmethod
    def fetch_all_workspaces_from_drive(
            cls,
            excel_key: dict,
            delegate_user: str,
            folder_id: str,
            max_workers: int = 20
    ) -> dict:
        if not HAS_GOOGLE_DRIVE:
            raise ImportError("Google Drive client libraries missing. Install google-api-python-client google-auth")

        scopes = ['https://www.googleapis.com/auth/drive']
        credentials = service_account.Credentials.from_service_account_info(excel_key, scopes=scopes)
        if delegate_user:
            credentials = credentials.with_subject(delegate_user)

        drive_service = build('drive', 'v3', credentials=credentials)
        existing_files = cls.list_workspace_files_in_folder(excel_key, delegate_user, folder_id)

        all_workspace_configs = {}
        with ThreadPoolExecutor(max_workers=max_workers) as executor:
            future_to_ws = {
                executor.submit(cls.fetch_and_parse_from_drive, drive_service, item["file_id"]): ws_id
                for ws_id, item in existing_files.items()
            }
            for future in future_to_ws:
                ws_id = future_to_ws[future]
                try:
                    all_workspace_configs[ws_id] = future.result()
                except Exception as e:
                    print(f"Error fetching workspace {ws_id}: {e}")

        return all_workspace_configs

    @classmethod
    def create_workspace_config_if_not_exists(
            cls,
            excel_key: dict,
            delegate_user: str,
            folder_id: str,
            workspace_id: str | int,
            market_name: str,
            configs_dict: dict = None
    ) -> str | None:
        if not HAS_GOOGLE_DRIVE:
            raise ImportError("Google Drive client libraries missing. Install google-api-python-client google-auth")

        scopes = [
            'https://www.googleapis.com/auth/drive',
            'https://www.googleapis.com/auth/spreadsheets'
        ]
        credentials = service_account.Credentials.from_service_account_info(excel_key, scopes=scopes)
        if delegate_user:
            credentials = credentials.with_subject(delegate_user)

        drive_service = build('drive', 'v3', credentials=credentials)

        existing_files = cls.list_workspace_files_in_folder(excel_key, delegate_user, folder_id)
        str_ws_id = str(workspace_id)

        if str_ws_id in existing_files:
            existing_file_id = existing_files[str_ws_id]["file_id"]
            print(
                f"Spreadsheet for workspace [{workspace_id}] already exists (ID: {existing_file_id}). Skipping creation.")
            return existing_file_id

        wb = cls.create_excel_workbook(configs_dict)
        excel_buffer = io.BytesIO()
        wb.save(excel_buffer)
        excel_buffer.seek(0)

        target_title = f"{market_name} [{workspace_id}]"

        media = MediaIoBaseUpload(
            excel_buffer,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            resumable=True
        )

        file_metadata = {
            'name': target_title,
            'parents': [folder_id],
            'mimeType': 'application/vnd.google-apps.spreadsheet'
        }

        file = drive_service.files().create(
            body=file_metadata,
            media_body=media,
            supportsAllDrives=True,
            fields='id'
        ).execute()

        new_file_id = file.get('id')
        print(f"Successfully created new spreadsheet for workspace [{workspace_id}] (ID: {new_file_id}).")
        return new_file_id